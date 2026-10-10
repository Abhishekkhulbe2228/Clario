from pathlib import Path
from typing import Iterable
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument


SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def load_file(path: Path, source_name: str | None = None) -> list[Document]:
    """
    Extracts text documents from a supported file format.
    Accepts an optional source_name to ensure document metadata reflects
    the original filename rather than temporary scratch paths.
    """
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        docs = PyPDFLoader(str(path)).load()
    elif suffix in {".txt", ".md"}:
        docs = TextLoader(str(path), encoding="utf-8").load()
    elif suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        docs = [Document(page_content=text, metadata={"source": source_name or str(path)})]
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    if source_name:
        for d in docs:
            d.metadata["source"] = source_name
    return docs


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)
    return splitter.split_documents(list(docs))